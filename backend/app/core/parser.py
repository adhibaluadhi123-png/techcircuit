"""Document Ingestion Parser for CANON Brand Governance.
Validates file signatures (magic bytes) to prevent deceptive file uploads,
enforces defensive format parsing, and handles text, documents, and image assets.
"""
import io
import re
from typing import Tuple, Dict, Any, Optional
import docx
import pypdf
from app.core.visual_audit import analyze_image_asset

# Magic Byte Signatures for supported file formats
FILE_SIGNATURES = {
    "pdf": [b"%PDF-"],
    "docx": [b"PK\x03\x04"],
    "png": [b"\x89PNG\r\n\x1a\n"],
    "jpg": [b"\xff\xd8\xff"],
    "jpeg": [b"\xff\xd8\xff"],
}


def validate_file_signature(filename: str, content_bytes: bytes) -> str:
    """
    Validate that the file header magic bytes match the stated file extension.
    Rejects deceptive files (e.g. executable or script renamed to .pdf).
    """
    if not content_bytes or len(content_bytes) == 0:
        raise ValueError("Uploaded file is empty (0 bytes).")

    ext = filename.lower().split(".")[-1] if "." in filename else ""
    if not ext:
        raise ValueError("Missing file extension.")

    # Check for known dangerous executable headers
    if content_bytes.startswith(b"MZ") or content_bytes.startswith(b"\x7fELF") or content_bytes.startswith(b"#!"):
        raise ValueError("Dangerous file signature detected: binary executables and scripts are prohibited.")

    if ext in FILE_SIGNATURES:
        valid_signatures = FILE_SIGNATURES[ext]
        matches = any(content_bytes.startswith(sig) for sig in valid_signatures)
        if not matches:
            raise ValueError(f"File signature mismatch for '.{ext}': deceptive or corrupted file format.")
    elif ext == "txt":
        # Plain text must not contain binary null bytes
        if b"\x00" in content_bytes[:4096]:
            raise ValueError("Binary content detected in plain text file.")
    else:
        raise ValueError(f"Unsupported file extension '.{ext}'. Supported: .pdf, .docx, .txt, .png, .jpg, .jpeg")

    return ext


def extract_text_from_txt(content_bytes: bytes) -> str:
    """Extract and decode plain text."""
    try:
        return content_bytes.decode("utf-8")
    except UnicodeDecodeError:
        return content_bytes.decode("latin-1", errors="replace")


def extract_text_from_docx(content_bytes: bytes) -> str:
    """Extract paragraphs and table text from DOCX with malformed file handling."""
    try:
        doc = docx.Document(io.BytesIO(content_bytes))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    paragraphs.append(" | ".join(row_texts))
        return "\n\n".join(paragraphs)
    except Exception as e:
        raise ValueError(f"Malformed or corrupted DOCX document: {str(e)}")


def extract_text_from_pdf(content_bytes: bytes) -> str:
    """Extract text from PDF pages using pypdf with malformed file handling."""
    try:
        reader = pypdf.PdfReader(io.BytesIO(content_bytes))
        pages_text = []
        for idx, page in enumerate(reader.pages):
            try:
                extracted = page.extract_text()
                if extracted:
                    pages_text.append(extracted.strip())
            except Exception:
                continue
        if not pages_text:
            return "[PDF scanned or empty text layer. Visual inspection required.]"
        return "\n\n".join(pages_text)
    except Exception as e:
        raise ValueError(f"Malformed or corrupted PDF document: {str(e)}")


def parse_uploaded_file(filename: str, content_bytes: bytes) -> str:
    """
    Dispatcher based on validated file format.
    Validates file magic bytes and extracts clean content.
    """
    ext = validate_file_signature(filename, content_bytes)

    if ext == "txt":
        return extract_text_from_txt(content_bytes)
    elif ext == "docx":
        return extract_text_from_docx(content_bytes)
    elif ext == "pdf":
        return extract_text_from_pdf(content_bytes)
    elif ext in ["png", "jpg", "jpeg"]:
        analysis = analyze_image_asset(content_bytes, filename)
        dominant_colors_str = ", ".join(c["hex"] for c in analysis["dominant_colors"][:3])
        return (
            f"[Image Asset: {filename} ({analysis['width']}x{analysis['height']}px, format: {analysis['format']}). "
            f"Aspect ratio: {analysis['aspect_ratio']}:1. Dominant palette: {dominant_colors_str}. "
            f"Visual brand and contrast checks completed.]"
        )
    else:
        raise ValueError(f"Unsupported file type '.{ext}'. Supported: .pdf, .docx, .txt, .png, .jpg, .jpeg")


async def read_uploaded_file_safely(file: Any, max_size_mb: int) -> bytes:
    """
    Read an uploaded file in 64KB stream chunks to enforce upload size limits
    without loading oversized files into memory. Validates non-empty content.
    """
    chunk_size = 64 * 1024
    max_bytes = max_size_mb * 1024 * 1024
    chunks = []
    total_read = 0

    while True:
        chunk = await file.read(chunk_size)
        if not chunk:
            break
        total_read += len(chunk)
        if total_read > max_bytes:
            raise ValueError(f"Uploaded file exceeds maximum permitted size of {max_size_mb}MB.")
        chunks.append(chunk)

    data = b"".join(chunks)
    if len(data) == 0:
        raise ValueError("Uploaded file is empty (0 bytes).")
    return data

